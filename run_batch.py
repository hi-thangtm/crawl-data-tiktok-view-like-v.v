import asyncio
from collections import Counter
from datetime import datetime
from pathlib import Path
from time import perf_counter

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Font,
    PatternFill,
)
from openpyxl.utils import get_column_letter

from app.collectors.tiktok_batch import TikTokBatchCollector
from app.utils.url_parser import parse_tiktok_url


# ============================================================
# CONFIG
# ============================================================

APP_NAME = "TikTok Video Performance"

# Khi chay bang Python trong thu muc project, file urls.txt van duoc
# doc tu dung thu muc data cua project nhu truoc day.
PROJECT_DIR = Path(__file__).resolve().parent
INPUT_FILE = PROJECT_DIR / "data" / "urls.txt"

# Excel KHONG luu ben trong project / .app nua.
# Moi nguoi dung se co thu muc export rieng trong Documents.
PREFERRED_EXPORT_DIR = (
    Path.home()
    / "Documents"
    / APP_NAME
    / "Exports"
)

# Neu macOS / may nguoi dung khong cho ghi vao Documents,
# tool se fallback ve thu muc trong Home thay vi bi crash.
FALLBACK_EXPORT_DIR = (
    Path.home()
    / APP_NAME
    / "Exports"
)

# Uu tien on dinh hon toc do.
CONCURRENCY = 3

RETRY_CONCURRENCY = 1


# ============================================================
# EXPORT DIRECTORY
# ============================================================

def get_export_dir():

    candidates = [
        PREFERRED_EXPORT_DIR,
        FALLBACK_EXPORT_DIR,
    ]

    last_error = None

    for export_dir in candidates:

        try:

            export_dir.mkdir(
                parents=True,
                exist_ok=True,
            )

            # Kiem tra thu muc thuc su co the ghi duoc.
            test_file = (
                export_dir
                / ".write_test"
            )

            test_file.write_text(
                "ok",
                encoding="utf-8",
            )

            try:
                test_file.unlink()
            except OSError:
                pass

            return export_dir

        except OSError as exc:

            last_error = exc

    raise RuntimeError(
        "Khong tao duoc thu muc xuat Excel. "
        f"Loi cuoi: {last_error}"
    )


# ============================================================
# STATUS HIEN THI
# ============================================================

STATUS_LABELS = {
    "THANH_CONG": "Thành công",

    "TAM_BI_GIOI_HAN_403":
        "TikTok tạm giới hạn truy cập",

    "VIDEO_KHONG_KHA_DUNG":
        "Video không khả dụng",

    "VIDEO_RIENG_TU":
        "Video riêng tư",

    "ACCESS_CHALLENGE":
        "TikTok yêu cầu xác minh",

    "SESSION_HET_HAN":
        "TikTok Session hết hạn",

    "CAN_RETRY":
        "Cần thử lại",

    "CAN_RETRY_403":
        "Cần thử lại (403)",

    "TIMEOUT":
        "Timeout",

    "KHONG_HO_TRO_NEN_TANG":
        "Không hỗ trợ nền tảng",

    "LINK_TIKTOK_KHONG_HOP_LE":
        "Link TikTok không hợp lệ",

    "LINK_TRONG":
        "Link trống",
}


# ============================================================
# LOAD URL
# ============================================================

def load_urls():

    if not INPUT_FILE.exists():

        raise FileNotFoundError(
            f"Khong tim thay file: {INPUT_FILE}"
        )

    lines = INPUT_FILE.read_text(
        encoding="utf-8"
    ).splitlines()

    raw_urls = []

    for line in lines:

        url = line.strip()

        if url:
            raw_urls.append(url)

    return raw_urls


# ============================================================
# LOC LINK TRUNG
# ============================================================

def deduplicate_urls(
    urls,
):

    unique_urls = []

    duplicate_urls = []

    seen = set()

    for url in urls:

        parsed = parse_tiktok_url(
            url
        )

        # Link TikTok hop le:
        # dung normalized URL de phat hien trung.
        if (
            parsed.valid
            and parsed.normalized_url
        ):

            key = parsed.normalized_url

        else:

            # Link khong hop le:
            # van giu lai de collector phan loai.
            key = url.strip()

        if key in seen:

            duplicate_urls.append(
                url
            )

            continue

        seen.add(key)

        unique_urls.append(
            url
        )

    return (
        unique_urls,
        duplicate_urls,
    )


# ============================================================
# FORMAT STATUS
# ============================================================

def status_label(
    status,
):

    if status in STATUS_LABELS:

        return STATUS_LABELS[
            status
        ]

    if status.startswith(
        "HTTP_"
    ):

        return (
            f"Lỗi HTTP "
            f"{status.replace('HTTP_', '')}"
        )

    if status.startswith(
        "LOI:"
    ):

        return (
            "Lỗi collector: "
            + status.replace(
                "LOI:",
                "",
            )
        )

    return status


# ============================================================
# TERMINAL SUMMARY
# ============================================================

def print_summary(
    results,
    duplicate_count,
    original_count,
    elapsed,
):

    counts = Counter(
        result.status
        for result in results
    )

    success = counts[
        "THANH_CONG"
    ]

    unavailable = counts[
        "VIDEO_KHONG_KHA_DUNG"
    ]

    private = counts[
        "VIDEO_RIENG_TU"
    ]

    challenge = counts[
        "ACCESS_CHALLENGE"
    ]

    session_expired = counts[
        "SESSION_HET_HAN"
    ]

    retry = (
        counts["CAN_RETRY"]
        + counts["CAN_RETRY_403"]
        + counts["TAM_BI_GIOI_HAN_403"]
        + counts["TIMEOUT"]
    )

    known_statuses = {
        "THANH_CONG",
        "VIDEO_KHONG_KHA_DUNG",
        "VIDEO_RIENG_TU",
        "ACCESS_CHALLENGE",
        "SESSION_HET_HAN",
        "CAN_RETRY",
        "CAN_RETRY_403",
        "TAM_BI_GIOI_HAN_403",
        "TIMEOUT",
    }

    errors = sum(
        count
        for status, count
        in counts.items()
        if status not in known_statuses
    )

    print()
    print(
        "=" * 80
    )

    print(
        "TONG KET"
    )

    print(
        "=" * 80
    )

    print(
        f"URL ban dau               : "
        f"{original_count}"
    )

    print(
        f"Link trung                : "
        f"{duplicate_count}"
    )

    print(
        f"Link duoc xu ly           : "
        f"{len(results)}"
    )

    print()
    print(
        f"Lay duoc du lieu          : "
        f"{success}"
    )

    print(
        f"Video khong kha dung      : "
        f"{unavailable}"
    )

    print(
        f"Video rieng tu            : "
        f"{private}"
    )

    print(
        f"Can xac minh              : "
        f"{challenge}"
    )

    print(
        f"Session het han           : "
        f"{session_expired}"
    )

    print(
        f"Can thu lai               : "
        f"{retry}"
    )

    print(
        f"Link / loi khac           : "
        f"{errors}"
    )

    print()
    print(
        f"Tong thoi gian            : "
        f"{elapsed:.2f} giay"
    )

    if results:

        print(
            f"Quy doi trung binh        : "
            f"{elapsed / len(results):.2f} "
            f"giay/link"
        )

    print(
        "=" * 80
    )


# ============================================================
# EXCEL STYLE
# ============================================================

HEADER_FILL = PatternFill(
    "solid",
    fgColor="0F766E",
)

HEADER_FONT = Font(
    color="FFFFFF",
    bold=True,
)

SUCCESS_FILL = PatternFill(
    "solid",
    fgColor="DCFCE7",
)

WARNING_FILL = PatternFill(
    "solid",
    fgColor="FEF3C7",
)

ERROR_FILL = PatternFill(
    "solid",
    fgColor="FEE2E2",
)

STATIC_FILL = PatternFill(
    "solid",
    fgColor="F3F4F6",
)


# ============================================================
# EXPORT EXCEL
# ============================================================

def export_excel(
    results,
    duplicate_urls,
    original_count,
    elapsed,
):

    export_dir = get_export_dir()

    timestamp = datetime.now().strftime(
        "%Y%m%d_%H%M%S"
    )

    output_file = (
        export_dir
        / f"tiktok_results_{timestamp}.xlsx"
    )

    wb = Workbook()

    ws = wb.active

    ws.title = "Ket qua"

    ws.sheet_view.showGridLines = False

    headers = [
        "STT",
        "Link video",
        "Title / Caption",
        "Uploader / Creator",
        "Username",
        "Thời gian đăng",
        "View",
        "Like",
        "Comment",
        "Share",
        "Save / Collection",
        "Thời gian lấy dữ liệu",
        "Trạng thái",
        "Engagement",
        "Engagement + Save",
        "Collector",
    ]

    for col_index, header in enumerate(
        headers,
        start=1,
    ):

        cell = ws.cell(
            row=1,
            column=col_index,
            value=header,
        )

        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    # ========================================================
    # DATA ROWS
    # ========================================================

    for index, result in enumerate(
        results,
        start=1,
    ):

        row = index + 1

        values = [
            index,
            result.normalized_url
            or result.link,
            result.title,
            result.uploader,
            result.username,
            result.release_time,
            result.view,
            result.like,
            result.comment,
            result.share,
            result.save,
            result.retrieval_time,
            status_label(
                result.status
            ),
        ]

        for col_index, value in enumerate(
            values,
            start=1,
        ):

            ws.cell(
                row=row,
                column=col_index,
                value=value,
            )

        # ====================================================
        # ENGAGEMENT
        #
        # Derived value -> Excel formula
        # ====================================================

        ws.cell(
            row=row,
            column=14,
            value=(
                f'=IF(G{row}>0,'
                f'(H{row}+I{row}+J{row})'
                f'/G{row},"")'
            ),
        )

        # Engagement + Save
        ws.cell(
            row=row,
            column=15,
            value=(
                f'=IF(G{row}>0,'
                f'(H{row}+I{row}+J{row}+K{row})'
                f'/G{row},"")'
            ),
        )

        ws.cell(
            row=row,
            column=16,
            value=result.collector,
        )

        # Percentage format
        ws.cell(
            row=row,
            column=14,
        ).number_format = "0.00%"

        ws.cell(
            row=row,
            column=15,
        ).number_format = "0.00%"

        # Number format
        for col in range(
            7,
            12,
        ):

            ws.cell(
                row=row,
                column=col,
            ).number_format = "#,##0"

        # ====================================================
        # STATUS COLOR
        # ====================================================

        status_cell = ws.cell(
            row=row,
            column=13,
        )

        if result.status == "THANH_CONG":

            status_cell.fill = (
                SUCCESS_FILL
            )

        elif result.status in {
            "VIDEO_KHONG_KHA_DUNG",
            "VIDEO_RIENG_TU",
        }:

            status_cell.fill = (
                STATIC_FILL
            )

        elif result.status in {
            "CAN_RETRY",
            "CAN_RETRY_403",
            "TAM_BI_GIOI_HAN_403",
            "TIMEOUT",
            "ACCESS_CHALLENGE",
        }:

            status_cell.fill = (
                WARNING_FILL
            )

        else:

            status_cell.fill = (
                ERROR_FILL
            )

    # ========================================================
    # ALIGNMENT
    # ========================================================

    for row in ws.iter_rows(
        min_row=2,
    ):

        for cell in row:

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

    # ========================================================
    # FREEZE + FILTER
    # ========================================================

    ws.freeze_panes = "A2"

    ws.auto_filter.ref = (
        f"A1:P{ws.max_row}"
    )

    # ========================================================
    # COLUMN WIDTH
    # ========================================================

    widths = {
        "A": 7,
        "B": 52,
        "C": 55,
        "D": 25,
        "E": 24,
        "F": 21,
        "G": 15,
        "H": 15,
        "I": 15,
        "J": 15,
        "K": 18,
        "L": 21,
        "M": 27,
        "N": 16,
        "O": 18,
        "P": 26,
    }

    for column, width in widths.items():

        ws.column_dimensions[
            column
        ].width = width

    ws.row_dimensions[1].height = 32

    # ========================================================
    # SUMMARY SHEET
    # ========================================================

    summary = wb.create_sheet(
        "Tong quan"
    )

    summary.sheet_view.showGridLines = False

    counts = Counter(
        result.status
        for result in results
    )

    summary_data = [
        (
            "CHỈ SỐ",
            "GIÁ TRỊ",
        ),
        (
            "URL ban đầu",
            original_count,
        ),
        (
            "Link trùng đã loại",
            len(duplicate_urls),
        ),
        (
            "Link thực tế xử lý",
            len(results),
        ),
        (
            "Lấy được dữ liệu",
            counts["THANH_CONG"],
        ),
        (
            "Video không khả dụng",
            counts[
                "VIDEO_KHONG_KHA_DUNG"
            ],
        ),
        (
            "Video riêng tư",
            counts[
                "VIDEO_RIENG_TU"
            ],
        ),
        (
            "TikTok yêu cầu xác minh",
            counts[
                "ACCESS_CHALLENGE"
            ],
        ),
        (
            "TikTok Session hết hạn",
            counts[
                "SESSION_HET_HAN"
            ],
        ),
        (
            "Cần thử lại",
            (
                counts["CAN_RETRY"]
                + counts[
                    "CAN_RETRY_403"
                ]
                + counts[
                    "TAM_BI_GIOI_HAN_403"
                ]
                + counts["TIMEOUT"]
            ),
        ),
        (
            "Tổng thời gian (giây)",
            round(
                elapsed,
                2,
            ),
        ),
    ]

    for row_index, row_data in enumerate(
        summary_data,
        start=1,
    ):

        for col_index, value in enumerate(
            row_data,
            start=1,
        ):

            cell = summary.cell(
                row=row_index,
                column=col_index,
                value=value,
            )

            if row_index == 1:

                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT

    summary.column_dimensions[
        "A"
    ].width = 35

    summary.column_dimensions[
        "B"
    ].width = 20

    # ========================================================
    # DUPLICATE SHEET
    # ========================================================

    if duplicate_urls:

        duplicate_sheet = (
            wb.create_sheet(
                "Link trung"
            )
        )

        duplicate_sheet[
            "A1"
        ] = "Link bị trùng"

        duplicate_sheet[
            "A1"
        ].fill = HEADER_FILL

        duplicate_sheet[
            "A1"
        ].font = HEADER_FONT

        for row_index, url in enumerate(
            duplicate_urls,
            start=2,
        ):

            duplicate_sheet.cell(
                row=row_index,
                column=1,
                value=url,
            )

        duplicate_sheet.column_dimensions[
            "A"
        ].width = 100

    # ========================================================
    # EXCEL RECALC
    # ========================================================

    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"

    wb.save(
        output_file
    )

    return output_file


# ============================================================
# MAIN
# ============================================================

async def main():

    print()
    print(
        "=" * 80
    )

    print(
        "TIKTOK VIDEO PERFORMANCE TOOL"
    )

    print(
        "=" * 80
    )

    # ========================================================
    # LOAD
    # ========================================================

    raw_urls = load_urls()

    print()
    print(
        "URL trong file:",
        len(raw_urls),
    )

    # ========================================================
    # DEDUPE
    # ========================================================

    (
        urls,
        duplicate_urls,
    ) = deduplicate_urls(
        raw_urls
    )

    print(
        "Link trung:",
        len(duplicate_urls),
    )

    print(
        "Link se xu ly:",
        len(urls),
    )

    if not urls:

        print()
        print(
            "Khong co URL de xu ly."
        )

        return

    # ========================================================
    # COLLECT
    # ========================================================

    collector = (
        TikTokBatchCollector(
            concurrency=CONCURRENCY,
            retry_concurrency=1,
            start_interval=0.45,
            retry_start_interval=1.0,
            cooldown_403=12.0,
            retry_delay=6.0,
            headless=True,
        )
    )

    start = perf_counter()

    results = (
        await collector.collect_many(
            urls
        )
    )

    elapsed = (
        perf_counter()
        - start
    )

    # ========================================================
    # SUMMARY
    # ========================================================

    print_summary(
        results=results,
        duplicate_count=(
            len(
                duplicate_urls
            )
        ),
        original_count=(
            len(
                raw_urls
            )
        ),
        elapsed=elapsed,
    )

    # ========================================================
    # EXPORT
    # ========================================================

    output_file = export_excel(
        results=results,
        duplicate_urls=(
            duplicate_urls
        ),
        original_count=(
            len(
                raw_urls
            )
        ),
        elapsed=elapsed,
    )

    print()
    print(
        "DA XUAT EXCEL:"
    )

    print(
        output_file.resolve()
    )

    print()


# ============================================================
# RUN
# ============================================================

if __name__ == "__main__":

    asyncio.run(
        main()
    )