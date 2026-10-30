from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Font,
    PatternFill,
    Border,
    Side,
)
from openpyxl.utils import get_column_letter

from app.models import TikTokProfileResult
from run_batch import get_export_dir


HEADER_FILL = PatternFill(
    "solid",
    fgColor="0F766E",
)

HEADER_FONT = Font(
    color="FFFFFF",
    bold=True,
    size=11,
)

SUCCESS_FILL = PatternFill(
    "solid",
    fgColor="DCFCE7",
)

WARNING_FILL = PatternFill(
    "solid",
    fgColor="FEF3C7",
)

ERROR_FILL = PatternFill(
    "solid",
    fgColor="FEE2E2",
)

STATIC_FILL = PatternFill(
    "solid",
    fgColor="F3F4F6",
)

THIN_BORDER = Border(
    left=Side(style="thin", color="E5E7EB"),
    right=Side(style="thin", color="E5E7EB"),
    top=Side(style="thin", color="E5E7EB"),
    bottom=Side(style="thin", color="E5E7EB"),
)


PROFILE_STATUS_LABELS = {
    "SUCCESS": "Thành công",
    "PARTIAL": "Một phần dữ liệu",
    "NOT_FOUND": "Không tìm thấy kênh",
    "PRIVATE": "Tài khoản riêng tư",
    "CAPTCHA_REQUIRED": "Cần xác minh Captcha",
    "RATE_LIMITED": "Bị giới hạn (403/429)",
    "SESSION_EXPIRED": "Session hết hạn",
    "TIMEOUT": "Quá thời gian (Timeout)",
    "NETWORK_ERROR": "Lỗi kết nối mạng",
    "PARSE_ERROR": "Lỗi phân tích dữ liệu",
    "LINK_TIKTOK_KHONG_HOP_LE": "Link/Username không hợp lệ",
    "KHONG_HO_TRO_NEN_TANG": "Không hỗ trợ nền tảng",
    "LINK_TRONG": "Link trống",
    "CANCELLED": "Đã dừng bởi người dùng",
}


def sanitize_text(val: Any) -> Any:
    """
    Tranh CSV/Excel formula injection.
    Neu chuoi bat dau bang '=', '+', '-', '@', '\t', '\r' thi them dau nha' don (').
    """
    if val is None:
        return ""
    if isinstance(val, (int, float, bool)):
        return val
    s = str(val)
    if not s:
        return ""
    if s[0] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + s
    return s


def format_profile_status(status: str) -> str:
    return PROFILE_STATUS_LABELS.get(status, status)


def export_profile_excel(
    results: list[TikTokProfileResult],
    duplicate_count: int = 0,
    original_count: int = 0,
    elapsed: float = 0.0,
) -> Path:
    """
    Xuat danh sach ho so kenh TikTok ra tap tin Excel dep mat:
    - Thu muc: ~/Documents/TikTok Video Performance/Exports/
    - Ten file: TikTok_Profiles_YYYYMMDD_HHMMSS.xlsx
    - Freeze header row
    - AutoFilter
    - So luu dang so nguyen
    - user_id va sec_uid luu dang text
    - Bio co wrap text
    """
    export_dir = get_export_dir()
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_file = export_dir / f"TikTok_Profiles_{timestamp}.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "Kênh TikTok"
    ws.sheet_view.showGridLines = True

    headers = [
        "STT",
        "Username",
        "Tên kênh / Nickname",
        "ID kênh (User ID)",
        "Follower",
        "Following",
        "Likes",
        "Videos",
        "Tích xanh (Verified)",
        "Tài khoản riêng tư",
        "Bio / Tiểu sử",
        "Website liên kết",
        "Link kênh",
        "Trạng thái",
        "Phương thức",
        "Thời gian thu thập",
        "Sec UID",
        "Độ chính xác số liệu",
        "Ghi chú / Lỗi",
        "Đầu vào gốc",
    ]

    ws.row_dimensions[1].height = 28

    for col_index, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_index, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # ========================================================
    # ROWS
    # ========================================================
    for index, res in enumerate(results, start=1):
        row = index + 1
        ws.row_dimensions[row].height = 24

        verified_str = "Có" if res.verified is True else ("Không" if res.verified is False else "-")
        private_str = "Có" if res.private_account is True else ("Không" if res.private_account is False else "-")

        values = [
            index,
            sanitize_text(res.username),
            sanitize_text(res.nickname),
            str(res.user_id) if res.user_id is not None else "",
            res.follower_count,
            res.following_count,
            res.likes_count,
            res.video_count,
            verified_str,
            private_str,
            sanitize_text(res.bio),
            sanitize_text(res.bio_website_url),
            sanitize_text(res.profile_url),
            format_profile_status(res.status),
            sanitize_text(res.source_method),
            sanitize_text(res.collected_at),
            str(res.sec_uid) if res.sec_uid is not None else "",
            sanitize_text(res.stats_precision),
            sanitize_text(res.error_message),
            sanitize_text(res.input_url),
        ]

        for col_index, val in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col_index)
            cell.border = THIN_BORDER

            # Kieu so nguyen
            if isinstance(val, int):
                cell.value = val
                cell.number_format = "#,##0"
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif val is None or val == "":
                cell.value = "-"
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.value = str(val)
                cell.data_type = "s"
                # user_id (col 4) va sec_uid (col 17) luon format dang text de tranh mat so
                if col_index in {4, 17}:
                    cell.number_format = "@"
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                # Bio (col 11) bat wrap text
                elif col_index == 11:
                    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                # Cac cot can giua
                elif col_index in {1, 9, 10, 14, 15, 16, 18}:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

            # To mau trang thai cot 14 (Trang thai)
            if col_index == 14:
                status = res.status
                if status == "SUCCESS":
                    cell.fill = SUCCESS_FILL
                elif status in {"PARTIAL", "PRIVATE"}:
                    cell.fill = WARNING_FILL
                elif status in {
                    "NOT_FOUND",
                    "CAPTCHA_REQUIRED",
                    "RATE_LIMITED",
                    "SESSION_EXPIRED",
                    "TIMEOUT",
                    "NETWORK_ERROR",
                    "PARSE_ERROR",
                    "LINK_TIKTOK_KHONG_HOP_LE",
                    "KHONG_HO_TRO_NEN_TANG",
                }:
                    cell.fill = ERROR_FILL
                else:
                    cell.fill = STATIC_FILL

    # ========================================================
    # FREEZE & AUTOFILTER
    # ========================================================
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # ========================================================
    # COLUMN WIDTHS
    # ========================================================
    base_widths = {
        1: 6,    # STT
        2: 20,   # Username
        3: 24,   # Tên kênh
        4: 22,   # ID kênh
        5: 14,   # Follower
        6: 12,   # Following
        7: 14,   # Likes
        8: 12,   # Videos
        9: 14,   # Verified
        10: 16,  # Private
        11: 36,  # Bio
        12: 24,  # Website
        13: 32,  # Link kênh
        14: 18,  # Trạng thái
        15: 16,  # Phương thức
        16: 20,  # Thời gian lấy
        17: 28,  # Sec UID
        18: 16,  # Độ chính xác
        19: 28,  # Ghi chú / Lỗi
        20: 30,  # Input gốc
    }

    for col_idx, width in base_widths.items():
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    wb.save(output_file)
    return output_file
